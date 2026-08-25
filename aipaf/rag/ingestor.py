"""
AIPAF — AI Project Assessment Framework
rag/ingestor.py — Ingestione documenti nel vector store ChromaDB

Responsabilità:
  - Legge PDF, DOCX, TXT e MD da docs/sources/regulatory/ e docs/sources/company/
  - Divide in chunk con overlap
  - Calcola checksum SHA-256 per rilevare modifiche
  - Indicizza in ChromaDB con metadati (source_type, filename, version, chunk_index)
  - Aggiorna docs/manifest.json con versioni e date di ingestione
  - Supporta re-ingestione selettiva: solo i file modificati vengono ri-processati

Embedding strategy:
  - Produzione  → OllamaEmbeddingFunction (riusa Ollama già presente)
                  Modello consigliato: nomic-embed-text (ollama pull nomic-embed-text)
  - Fallback    → HashEmbeddingFunction (deterministica, zero dipendenze,
                  usata anche nei test unitari)

Autore: Lorenzo Lombardi
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Optional

import chromadb
from chromadb.api.models.Collection import Collection
from chromadb.api.types import EmbeddingFunction

# ---------------------------------------------------------------------------
# Costanti e path di default
# ---------------------------------------------------------------------------

_PROJECT_ROOT  = Path(__file__).parent.parent.parent
_DOCS_ROOT     = _PROJECT_ROOT / "docs"
_SOURCES_ROOT  = _DOCS_ROOT / "sources"
_INDEX_PATH    = _DOCS_ROOT / "index"
_MANIFEST_PATH = _DOCS_ROOT / "manifest.json"

REGULATORY_DIR = _SOURCES_ROOT / "regulatory"
COMPANY_DIR    = _SOURCES_ROOT / "company"

# Parametri chunking
CHUNK_SIZE    = 800   # caratteri per chunk
CHUNK_OVERLAP = 150   # overlap tra chunk consecutivi


# ---------------------------------------------------------------------------
# Enumerazioni e dataclass
# ---------------------------------------------------------------------------

class DocumentSource(str, Enum):
    REGULATORY = "regulatory"   # EU AI Act, NIST AI RMF, ISO 42001, …
    COMPANY    = "company"      # Codice etico, policy interne, …


@dataclass
class IngestedDocument:
    """Rappresenta un documento processato con metadati."""
    filename:    str
    source_type: DocumentSource
    version:     str
    checksum:    str
    ingested_at: str
    chunk_count: int
    path:        Path = field(repr=False)


# ---------------------------------------------------------------------------
# Embedding functions
# ---------------------------------------------------------------------------

class HashEmbeddingFunction(EmbeddingFunction):
    """
    Embedding deterministico basato su SHA-256 del testo.

    Produce vettori a 128 dimensioni normalizzati in [-1, 1].
    Qualità semantica nulla, ma:
      - zero dipendenze esterne
      - zero download di modelli
      - completamente riproducibile → ideale per test unitari

    Estende chromadb.EmbeddingFunction: ChromaDB 1.5.x richiede name(),
    embed_query() e la configurazione serializzabile, altrimenti le query
    falliscono con AttributeError (v0.2.0 falliva in silenzio nel retriever).

    In produzione usa OllamaEmbeddingFunction.
    """

    DIM = 128

    def __init__(self, **kwargs):
        pass

    @staticmethod
    def name() -> str:
        return "hash_embedding_function"

    def get_config(self) -> dict:
        return {}

    @staticmethod
    def build_from_config(config: dict) -> "HashEmbeddingFunction":
        return HashEmbeddingFunction()

    @staticmethod
    def validate_config(config: dict) -> None:
        return None

    def validate_config_update(self, old_config: dict, new_config: dict) -> None:
        return None

    def __call__(self, input: list[str]) -> list[list[float]]:
        result = []
        for text in input:
            digest = hashlib.sha256(text.encode()).digest()
            # Estende il digest a DIM float in [-1, 1]
            vec: list[float] = []
            for i in range(self.DIM):
                byte_val = digest[i % len(digest)]
                vec.append((byte_val / 127.5) - 1.0)
            result.append(vec)
        return result


def get_embedding_function(
    backend: str = "ollama",
    ollama_url: str = "http://localhost:11434",
    ollama_model: str = "nomic-embed-text",
) -> object:
    """
    Restituisce la embedding function appropriata al backend.

    Args:
        backend:      "ollama" | "hash"
                      "hash" è il fallback deterministico (test / offline)
        ollama_url:   URL del server Ollama
        ollama_model: Modello embedding Ollama
                      Consigliato: nomic-embed-text (ollama pull nomic-embed-text)
                      Alternativa: mxbai-embed-large
    """
    if backend == "ollama":
        from chromadb.utils.embedding_functions import OllamaEmbeddingFunction
        return OllamaEmbeddingFunction(url=ollama_url, model_name=ollama_model)
    elif backend == "hash":
        return HashEmbeddingFunction()
    else:
        raise ValueError(f"Backend embedding non supportato: {backend!r}. Usa 'ollama' o 'hash'.")


# ---------------------------------------------------------------------------
# Estrattori di testo per tipo file
# ---------------------------------------------------------------------------

def _extract_text_pdf(path: Path) -> str:
    """Estrae testo grezzo da un PDF tramite pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)
    except ImportError:
        raise ImportError(
            "pypdf non installato. Eseguire: pip install pypdf"
        )
    except Exception as e:
        raise RuntimeError(f"Errore lettura PDF {path.name}: {e}") from e


def _extract_text_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_text_docx(path: Path) -> str:
    """
    Estrae testo da DOCX tramite python-docx, nell'ordine del documento,
    INCLUSE le tabelle (doc.paragraphs le ignora: nel framework AIPAF criteri,
    pesi e gate sono quasi tutti in tabella).
    """
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise ImportError(
            "python-docx non installato. Eseguire: pip install python-docx"
        )

    doc = docx.Document(str(path))
    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells: list[str] = []
                for cell in row.cells:
                    t = " ".join(cell.text.split())
                    if t and (not cells or cells[-1] != t):   # celle unite ripetono il testo
                        cells.append(t)
                if cells:
                    parts.append(" | ".join(cells))
            parts.append("")   # separatore dopo la tabella
    return "\n".join(parts)


_EXTRACTORS = {
    ".pdf":  _extract_text_pdf,
    ".txt":  _extract_text_txt,
    ".md":   _extract_text_txt,
    ".docx": _extract_text_docx,
}

SUPPORTED_EXTENSIONS = set(_EXTRACTORS.keys())


def extract_text(path: Path) -> str:
    """Dispatcher: estrae testo dal file in base all'estensione."""
    ext = path.suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(
            f"Formato non supportato: {ext}. "
            f"Supportati: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    return extractor(path)


# ---------------------------------------------------------------------------
# Chunker
# ---------------------------------------------------------------------------

def chunk_text(
    text: str,
    chunk_size:    int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
) -> list[str]:
    """
    Divide il testo in chunk con overlap.

    Strategia: rispetta i confini di paragrafo quando possibile.
    Se un paragrafo è più lungo di chunk_size, lo taglia comunque.
    """
    # Normalizza spazi multipli e righe vuote
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    text = re.sub(r" {2,}", " ", text)

    if len(text) <= chunk_size:
        return [text] if text.strip() else []

    chunks: list[str] = []
    start = 0

    while start < len(text):
        end = start + chunk_size

        if end >= len(text):
            chunk = text[start:].strip()
            if chunk:
                chunks.append(chunk)
            break

        # Cerca il punto di taglio migliore: fine paragrafo o fine frase
        cut = end
        for sep in ("\n\n", "\n", ". ", "? ", "! ", " "):
            idx = text.rfind(sep, start + chunk_overlap, end)
            if idx != -1:
                cut = idx + len(sep)
                break

        chunk = text[start:cut].strip()
        if chunk:
            chunks.append(chunk)

        start = cut - chunk_overlap

    return chunks


# ---------------------------------------------------------------------------
# Manifest management
# ---------------------------------------------------------------------------

def _load_manifest() -> dict:
    if _MANIFEST_PATH.exists():
        with open(_MANIFEST_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {"regulatory": {}, "company": {}}


def _save_manifest(manifest: dict) -> None:
    _MANIFEST_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(_MANIFEST_PATH, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False)


def _file_checksum(path: Path) -> str:
    """SHA-256 del file in hex."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def _is_already_ingested(manifest: dict, source_type: str, filename: str, checksum: str) -> bool:
    """True se il file è già indicizzato con lo stesso checksum."""
    entry = manifest.get(source_type, {}).get(filename)
    if entry is None:
        return False
    return entry.get("checksum") == checksum


# ---------------------------------------------------------------------------
# DocumentIngestor
# ---------------------------------------------------------------------------

class DocumentIngestor:
    """
    Indicizza documenti PDF/DOCX/TXT in ChromaDB.

    Usa due collection separate:
      - aipaf_regulatory  → documenti normativi (EU AI Act, NIST, ISO)
      - aipaf_company     → documenti aziendali (codice etico, policy)

    Usage:
        ingestor = DocumentIngestor()
        results = ingestor.ingest_all()          # indicizza tutto
        result  = ingestor.ingest_file(path, DocumentSource.REGULATORY)

        # Con Ollama embedding (produzione)
        ingestor = DocumentIngestor(embedding_backend="ollama")

        # Forza re-ingestione anche se il file non è cambiato
        ingestor.ingest_all(force=True)
    """

    COLLECTION_REGULATORY = "aipaf_regulatory"
    COLLECTION_COMPANY    = "aipaf_company"

    def __init__(
        self,
        index_path:        Path = _INDEX_PATH,
        embedding_backend: Optional[str] = "hash",
        ollama_url:        str  = "http://localhost:11434",
        ollama_model:      str  = "nomic-embed-text",
    ):
        """
        Args:
            index_path:        Directory dove ChromaDB persiste i dati.
            embedding_backend: "ollama" | "hash" | None
                               In produzione usa "ollama" con nomic-embed-text.
                               "hash" per test offline.
                               None = riusa la embedding function persistita
                               nell'indice (per status, query e assessment non
                               serve dichiararla di nuovo). Richiede un indice
                               gia' inizializzato.
            ollama_url:        URL Ollama (solo se backend="ollama").
            ollama_model:      Modello embedding Ollama.

        Raises:
            RuntimeError: backend diverso da quello con cui l'indice e' stato
                          creato (ChromaDB persiste la configurazione), oppure
                          backend None su indice vuoto.
        """
        index_path.mkdir(parents=True, exist_ok=True)
        self._client = chromadb.PersistentClient(path=str(index_path))
        self._manifest = _load_manifest()

        existing = {c.name for c in self._client.list_collections()}
        if embedding_backend is None:
            if not {self.COLLECTION_REGULATORY, self.COLLECTION_COMPANY} <= existing:
                raise RuntimeError(
                    f"Indice RAG non inizializzato in {index_path}: specificare il backend "
                    "embedding (ollama | hash) per crearlo, es. `aipaf rag ingest --backend ollama`."
                )
            self._ef = None   # ChromaDB ricostruisce la EF dalla configurazione persistita
        else:
            self._ef = get_embedding_function(embedding_backend, ollama_url, ollama_model)

        try:
            self._col_regulatory: Collection = self._get_or_create(
                self.COLLECTION_REGULATORY,
                "Documenti normativi AIPAF (EU AI Act, NIST, ISO 42001)",
            )
            self._col_company: Collection = self._get_or_create(
                self.COLLECTION_COMPANY,
                "Documenti aziendali AIPAF (codice etico, policy interne)",
            )
        except Exception as e:
            if "embedding function conflict" in str(e).lower():
                persisted = self.persisted_backend(index_path) or "?"
                raise RuntimeError(
                    f"L'indice in {index_path} e' stato creato con embedding '{persisted}', "
                    f"richiesto '{embedding_backend}'. Usare lo stesso backend, ometterlo "
                    "(riusa quello persistito) oppure eliminare la cartella dell'indice "
                    "e re-ingerire con `aipaf rag ingest --backend <nuovo> --force`."
                ) from e
            raise

    def _get_or_create(self, name: str, description: str) -> Collection:
        kwargs: dict = {"name": name, "metadata": {"description": description}}
        if self._ef is not None:
            kwargs["embedding_function"] = self._ef
        return self._client.get_or_create_collection(**kwargs)

    @staticmethod
    def persisted_backend(index_path: Path = _INDEX_PATH) -> Optional[str]:
        """Nome della embedding function persistita nell'indice (None se non inizializzato)."""
        try:
            client = chromadb.PersistentClient(path=str(index_path))
            col = client.get_collection(DocumentIngestor.COLLECTION_REGULATORY)
            cfg = getattr(col, "configuration_json", None) or {}
            name = (cfg.get("embedding_function") or {}).get("name")
            return {"ollama": "ollama", "hash_embedding_function": "hash"}.get(name, name)
        except Exception:
            return None

    # ------------------------------------------------------------------
    # API pubblica
    # ------------------------------------------------------------------

    def ingest_all(self, force: bool = False) -> list[IngestedDocument]:
        """
        Indicizza tutti i file in docs/sources/regulatory/ e docs/sources/company/.
        Salta i file già indicizzati con lo stesso checksum (a meno che force=True).

        Returns:
            Lista di IngestedDocument per i file effettivamente processati.
        """
        results: list[IngestedDocument] = []

        for source_dir, source_type in [
            (REGULATORY_DIR, DocumentSource.REGULATORY),
            (COMPANY_DIR,    DocumentSource.COMPANY),
        ]:
            if not source_dir.exists():
                source_dir.mkdir(parents=True, exist_ok=True)
                continue
            for path in sorted(source_dir.iterdir()):
                if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                    continue
                try:
                    doc = self.ingest_file(path, source_type, force=force)
                    if doc:
                        results.append(doc)
                        print(f"  ✓ [{source_type.value}] {path.name} — {doc.chunk_count} chunk")
                except Exception as e:
                    print(f"  ✗ [{source_type.value}] {path.name} — {e}")

        return results

    def ingest_file(
        self,
        path:        Path,
        source_type: DocumentSource,
        version:     str  = "",
        force:       bool = False,
    ) -> Optional[IngestedDocument]:
        """
        Indicizza un singolo file.

        Args:
            path:        Percorso assoluto al file.
            source_type: REGULATORY o COMPANY.
            version:     Stringa versione opzionale (es. "2024/1689").
                         Se vuota, usa data modifica del file.
            force:       Se True, re-indicizza anche se il checksum non è cambiato.

        Returns:
            IngestedDocument se il file è stato processato, None se saltato.
        """
        if not path.exists():
            raise FileNotFoundError(f"File non trovato: {path}")

        checksum = _file_checksum(path)
        filename = path.name

        if not force and _is_already_ingested(
            self._manifest, source_type.value, filename, checksum
        ):
            return None  # già aggiornato, skip

        # Estrai testo
        raw_text = extract_text(path)
        if not raw_text.strip():
            print(f"  ⚠ {filename}: testo vuoto, skip")
            return None

        chunks = chunk_text(raw_text)
        if not chunks:
            return None

        # Determina collection
        collection = (
            self._col_regulatory
            if source_type == DocumentSource.REGULATORY
            else self._col_company
        )

        # Rimuovi chunk precedenti di questo file (se esistono)
        self._delete_file_chunks(collection, filename)

        # Costruisci IDs e metadati
        doc_version = version or datetime.fromtimestamp(
            path.stat().st_mtime, tz=timezone.utc
        ).strftime("%Y-%m-%d")

        ids      = [f"{filename}::chunk::{i}" for i in range(len(chunks))]
        metadatas = [
            {
                "source_type":  source_type.value,
                "filename":     filename,
                "version":      doc_version,
                "checksum":     checksum,
                "chunk_index":  i,
                "chunk_total":  len(chunks),
            }
            for i in range(len(chunks))
        ]

        # Inserisci in ChromaDB
        collection.add(documents=chunks, metadatas=metadatas, ids=ids)

        # Aggiorna manifest
        ingested_at = datetime.now(tz=timezone.utc).isoformat()
        self._manifest.setdefault(source_type.value, {})[filename] = {
            "version":     doc_version,
            "checksum":    checksum,
            "ingested_at": ingested_at,
            "chunk_count": len(chunks),
        }
        _save_manifest(self._manifest)

        return IngestedDocument(
            filename=filename,
            source_type=source_type,
            version=doc_version,
            checksum=checksum,
            ingested_at=ingested_at,
            chunk_count=len(chunks),
            path=path,
        )

    def delete_file(self, filename: str, source_type: DocumentSource) -> bool:
        """
        Rimuove tutti i chunk di un file dall'indice e dal manifest.

        Returns:
            True se il file era presente e è stato rimosso, False se non trovato.
        """
        collection = (
            self._col_regulatory
            if source_type == DocumentSource.REGULATORY
            else self._col_company
        )
        deleted = self._delete_file_chunks(collection, filename)

        if filename in self._manifest.get(source_type.value, {}):
            del self._manifest[source_type.value][filename]
            _save_manifest(self._manifest)
            return True

        return deleted

    def status(self) -> dict:
        """
        Restituisce lo stato corrente dell'indice.

        Returns:
            Dict con conteggio documenti e chunk per tipo, più il manifest.
        """
        return {
            "regulatory": {
                "documents": len(self._manifest.get("regulatory", {})),
                "chunks":    self._col_regulatory.count(),
                "files":     list(self._manifest.get("regulatory", {}).keys()),
            },
            "company": {
                "documents": len(self._manifest.get("company", {})),
                "chunks":    self._col_company.count(),
                "files":     list(self._manifest.get("company", {}).keys()),
            },
        }

    # ------------------------------------------------------------------
    # Utility interne
    # ------------------------------------------------------------------

    def _delete_file_chunks(self, collection: Collection, filename: str) -> bool:
        """Elimina dalla collection tutti i chunk con metadato filename == filename."""
        try:
            existing = collection.get(where={"filename": filename})
            if existing["ids"]:
                collection.delete(ids=existing["ids"])
                return True
        except Exception:
            pass
        return False

    # Espone le collection per RAGRetriever
    @property
    def col_regulatory(self) -> Collection:
        return self._col_regulatory

    @property
    def col_company(self) -> Collection:
        return self._col_company
